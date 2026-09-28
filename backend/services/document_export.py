from io import BytesIO

from docx import Document
from docx.shared import Pt
from fpdf import FPDF


def clean_pdf_text(text: str) -> str:
    """Make text safe for FPDF built-in fonts."""

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u2026": "...",
        "\u00a0": " ",
        "\u2022": "-",
        "\u00b7": "-",
        "\u2011": "-",
        "\u2192": "->",
        "\u2190": "<-",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    # FPDF built-in Helvetica font works reliably with Latin-1.
    return text.encode("latin-1", "replace").decode("latin-1")


def format_pdf(
    text: str,
    doc_type: str = "Legal Document",
) -> bytes:
    """Create an A4 PDF from the generated legal document."""

    pdf = FPDF(
        orientation="P",
        unit="mm",
        format="A4",
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )

    pdf.set_margins(
        left=20,
        top=20,
        right=20,
    )

    pdf.add_page()

    # ---------------------------------------------------------
    # TITLE
    # ---------------------------------------------------------

    pdf.set_font(
        "Helvetica",
        "B",
        16,
    )

    title = clean_pdf_text(
        doc_type.upper()
    )

    pdf.multi_cell(
        0,
        10,
        title,
        align="C",
        border=0,
        wrapmode="CHAR",
    )

    pdf.ln(5)

    # ---------------------------------------------------------
    # BODY
    # ---------------------------------------------------------

    pdf.set_font(
        "Helvetica",
        "",
        11,
    )

    cleaned_text = clean_pdf_text(text)

    cleaned_text = cleaned_text.replace(
        "\r\n",
        "\n",
    )

    cleaned_text = cleaned_text.replace(
        "\r",
        "\n",
    )

    lines = cleaned_text.split("\n")

    heading_words = [
        "AGREEMENT",
        "CONTRACT",
        "PARTIES",
        "RECITALS",
        "DEFINITIONS",
        "RESPONSIBILITIES",
        "PAYMENT TERMS",
        "CONFIDENTIALITY",
        "TERMINATION",
        "GENERAL PROVISIONS",
        "SIGNATURES",
        "REVIEW NOTICE",
        "BACKGROUND",
        "SCOPE OF SERVICES",
        "USER-PROVIDED TERMS",
    ]

    for line in lines:

        line = line.strip()

        # Empty line
        if not line:
            pdf.ln(4)
            continue

        upper_line = line.upper()

        # Detect headings
        is_heading = any(
            upper_line == heading
            or upper_line.startswith(
                heading + ":"
            )
            or upper_line.startswith(
                heading + " "
            )
            for heading in heading_words
        )

        # Detect numbered sections
        if len(line) >= 3:
            if line[0].isdigit() and "." in line[:5]:
                is_heading = True

        # -----------------------------------------------------
        # HEADING
        # -----------------------------------------------------

        if is_heading:

            pdf.ln(3)

            pdf.set_font(
                "Helvetica",
                "B",
                12,
            )

            pdf.multi_cell(
                0,
                7,
                line,
                border=0,
                wrapmode="CHAR",
            )

            pdf.ln(1)

            pdf.set_font(
                "Helvetica",
                "",
                11,
            )

        # -----------------------------------------------------
        # NORMAL TEXT
        # -----------------------------------------------------

        else:

            pdf.multi_cell(
                0,
                6,
                line,
                border=0,
                wrapmode="CHAR",
            )

            pdf.ln(1)

    # ---------------------------------------------------------
    # FOOTER
    # ---------------------------------------------------------

    pdf.set_y(-15)

    pdf.set_font(
        "Helvetica",
        "I",
        8,
    )

    pdf.multi_cell(
        0,
        5,
        "LegalEase - AI-generated draft for human review",
        align="C",
        border=0,
        wrapmode="CHAR",
    )

    return bytes(
        pdf.output()
    )


def format_docx(
    text: str,
    doc_type: str = "Legal Document",
) -> bytes:
    """Create an editable DOCX document."""

    document = Document()

    # Title
    title = document.add_paragraph()

    title.alignment = 1

    run = title.add_run(
        doc_type.upper()
    )

    run.bold = True
    run.font.size = Pt(16)

    # Body
    for line in text.replace(
        "\r\n",
        "\n",
    ).split("\n"):

        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        run = paragraph.add_run(line)

        run.font.size = Pt(11)

    output = BytesIO()

    document.save(output)

    return output.getvalue()


def as_data(
    text: str,
    doc_type: str,
    format: str,
) -> bytes:
    """
    Convert generated legal text into the requested format.

    Supported formats:
    - txt
    - docx
    - pdf
    """

    file_format = format.lower().strip()

    if file_format == "txt":

        return text.encode(
            "utf-8"
        )

    if file_format == "docx":

        return format_docx(
            text=text,
            doc_type=doc_type,
        )

    if file_format == "pdf":

        return format_pdf(
            text=text,
            doc_type=doc_type,
        )

    raise ValueError(
        f"Unsupported export format: {format}"
    )


def make_export(
    text: str = "",
    doc_type: str = "Legal Document",
    file_format: str = "pdf",
    terms: str = "",
    parties: str = "",
    dates: str = "",
    **kwargs,
) -> bytes:

    if not text.strip():
        text_parts = []

        if parties.strip():
            text_parts.append("PARTIES\n" + parties.strip())

        if dates.strip():
            text_parts.append("DATES\n" + dates.strip())

        if terms.strip():
            text_parts.append("TERMS\n" + terms.strip())

        text = "\n\n".join(text_parts)

    selected_format = file_format.lower().strip()

    if selected_format == "txt":
        return text.encode("utf-8")

    if selected_format == "docx":
        return format_docx(
            text=text,
            doc_type=doc_type,
        )

    if selected_format == "pdf":
        return format_pdf(
            text=text,
            doc_type=doc_type,
        )

    raise ValueError(
        f"Unsupported export format: {file_format}"
    )